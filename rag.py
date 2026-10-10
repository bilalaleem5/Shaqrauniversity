"""Ingestion, retrieval (embeddings with BM25 fallback) and all LLM calls."""
import os, re, json, math, random, glob
from collections import Counter
import numpy as np
import config, logging
logging.getLogger('pypdf').setLevel(logging.ERROR)

_cache = {}
_client = None

# ---------------------------------------------------------------- helpers
def topic_name(path):
    n = os.path.splitext(os.path.basename(path))[0]
    n = re.sub(r"\s*\(\d+\)\s*", " ", n)      # drop "(1)", "(6)" copy suffixes
    return re.sub(r"\s+", " ", n).strip(" .")

def _norm(t):
    return re.sub(r"[ \t]+", " ", re.sub(r"\n{2,}", "\n", t or "")).strip()

def _pages(path):
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        from pypdf import PdfReader
        for i, p in enumerate(PdfReader(path).pages, 1):
            try: yield i, p.extract_text() or ""
            except Exception: yield i, ""
    elif ext == "docx":
        import docx
        d = docx.Document(path)
        yield 1, "\n".join(p.text for p in d.paragraphs)
    elif ext in ("xlsx", "xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for i, ws in enumerate(wb.worksheets, 1):
            rows = []
            for r in ws.iter_rows(values_only=True):
                cells = [str(c) for c in r if c is not None and str(c).strip()]
                if cells: rows.append(" | ".join(cells))
            yield i, f"Excel sheet '{ws.title}':\n" + "\n".join(rows)
    elif ext in ("txt", "md"):
        yield 1, open(path, encoding="utf-8", errors="ignore").read()

def _split(text, size=1000, overlap=150):
    text = _norm(text)
    if len(text) <= size: return [text] if len(text) > 40 else []
    out, i = [], 0
    while i < len(text):
        out.append(text[i:i + size]); i += size - overlap
    return out

# ---------------------------------------------------------------- ingestion
def paths(course):
    d = os.path.join(config.DATA_DIR, f"index_{course}")
    return d + ".json", d + ".npy"

def build_chunks(course):
    folder = config.COURSES[course]["folder"]
    chunks = []
    for f in sorted(glob.glob(os.path.join(folder, "*"))):
        if not f.lower().endswith((".pdf", ".docx", ".xlsx", ".xlsm", ".txt", ".md")): continue
        t = topic_name(f)
        for page, text in _pages(f):
            for piece in _split(text):
                chunks.append({"topic": t, "page": page, "text": piece})
    jp, np_ = paths(course)
    json.dump(chunks, open(jp, "w", encoding="utf-8"), ensure_ascii=False)
    if os.path.exists(np_): os.remove(np_)          # stale embeddings
    _cache.pop(course, None)
    return chunks

def _embed(texts):
    out = []
    for i in range(0, len(texts), 96):
        r = client().embeddings.create(model=config.EMBED_MODEL, input=texts[i:i + 96])
        out += [d.embedding for d in r.data]
    return np.array(out, dtype=np.float32)

def build_embeddings(course):
    chunks = load(course)["chunks"]
    m = _embed([c["text"] for c in chunks])
    m /= np.linalg.norm(m, axis=1, keepdims=True) + 1e-9
    np.save(paths(course)[1], m)
    _cache.pop(course, None)
    return len(chunks)

def load(course):
    if course in _cache: return _cache[course]
    jp, np_ = paths(course)
    if not os.path.exists(jp): build_chunks(course)
    chunks = json.load(open(jp, encoding="utf-8"))
    emb = np.load(np_) if os.path.exists(np_) else None
    if emb is not None and len(emb) != len(chunks): emb = None
    toks = [tokenize(c["text"]) for c in chunks]
    df = Counter(t for ts in toks for t in set(ts))
    _cache[course] = {"chunks": chunks, "emb": emb, "toks": toks, "df": df,
                      "avg": (sum(map(len, toks)) / max(1, len(toks)))}
    return _cache[course]

def topics(course):
    seen = []
    for c in load(course)["chunks"]:
        if c["topic"] not in seen: seen.append(c["topic"])
    return seen

# ---------------------------------------------------------------- retrieval
def tokenize(s): return re.findall(r"\w+", (s or "").lower())

def _bm25(ix, query, k):
    q = tokenize(query); N = len(ix["chunks"]); scores = []
    for toks in ix["toks"]:
        tf = Counter(toks); s = 0.0
        for w in q:
            if w not in tf: continue
            idf = math.log(1 + (N - ix["df"][w] + .5) / (ix["df"][w] + .5))
            s += idf * tf[w] * 2.2 / (tf[w] + 1.2 * (.25 + .75 * len(toks) / ix["avg"]))
        scores.append(s)
    order = np.argsort(scores)[::-1][:k]
    return [(i, scores[i]) for i in order if scores[i] > 0]

def retrieve(course, query, k=5):
    ix = load(course)
    if ix["emb"] is None and os.path.exists(paths(course)[1]):
        _cache.pop(course, None); ix = load(course)
    if ix["emb"] is not None and _have_llm():
        try:
            v = _embed([query])[0]; v /= np.linalg.norm(v) + 1e-9
            sims = ix["emb"] @ v
            order = np.argsort(sims)[::-1][:k]
            return [dict(ix["chunks"][i], score=float(sims[i])) for i in order]
        except Exception:
            pass
    return [dict(ix["chunks"][i], score=float(s)) for i, s in _bm25(ix, query, k)]

# ---------------------------------------------------------------- LLM
def _have_llm(): return bool(config.OPENAI_API_KEY) and not config.MOCK_LLM

def client():
    global _client
    if _client is None:
        from openai import OpenAI
        _client = OpenAI(api_key=config.OPENAI_API_KEY)
    return _client

def _chat(messages, json_mode=False, temperature=0.3, max_tokens=900):
    kw = {"response_format": {"type": "json_object"}} if json_mode else {}
    r = client().chat.completions.create(model=config.CHAT_MODEL, messages=messages,
                                         temperature=temperature, max_tokens=max_tokens, **kw)
    return r.choices[0].message.content.strip()

OOS = "[[OUT_OF_SCOPE]]"

def answer(course, question, history=None, intro=False):
    """Returns dict(answer, topic, in_scope, sources)."""
    cfg = config.COURSES[course]
    query = cfg["intro_query"] if intro else question
    ctx = retrieve(course, query, k=6 if intro else 5)
    topic = ctx[0]["topic"] if ctx else None
    sources = list(dict.fromkeys(f'{c["topic"]} (p.{c["page"]})' for c in ctx))[:4]
    if not _have_llm():
        if not ctx:
            return dict(answer="This question is outside the course material.", topic=None, in_scope=0, sources=[])
        return dict(answer="[Demo mode, no OpenAI key] Closest course material:\n\n" + ctx[0]["text"][:500],
                    topic=topic, in_scope=1, sources=sources)
    context = "\n\n".join(f'[{c["topic"]}, p.{c["page"]}]\n{c["text"]}' for c in ctx) or "(no relevant material found)"
    system = (
        cfg["style"] + "\n\nRULES:\n"
        "1. Answer ONLY using the COURSE MATERIAL below. Never use outside knowledge to introduce facts that are not "
        "supported by it. Simple worked examples and Excel/search steps that apply the material are allowed.\n"
        f"2. If the question cannot be answered from the course material (or is unrelated to the course), reply with "
        f"{OOS} followed by one short friendly sentence saying it is outside this course and suggesting what you can "
        "help with. Greetings and thanks are fine to answer normally.\n"
        "3. Reply in the same language the student used (English or Arabic). Keep answers clear, short and "
        "student-level; use short steps or bullet points when helpful.\n"
        "4. Mention which lecture/file the information comes from when it helps.\n"
        + ("5. The student asked for a short introduction to the course: give the course aim, main topics and how you "
           "can help (explain concepts, practice questions, quizzes). Offer to start with a first topic.\n" if intro else "")
        + f"\nCOURSE MATERIAL:\n{context}"
    )
    msgs = [{"role": "system", "content": system}]
    for h in (history or [])[-6:]:
        if h.get("role") in ("user", "assistant") and h.get("content"):
            msgs.append({"role": h["role"], "content": h["content"][:1500]})
    msgs.append({"role": "user", "content": question})
    text = _chat(msgs)
    in_scope = OOS not in text
    text = text.replace(OOS, "").strip()
    return dict(answer=text, topic=topic if in_scope else None, in_scope=int(in_scope),
                sources=sources if in_scope else [])

def make_question(course, topic, avoid=()):
    ix = load(course)
    pool = [c for c in ix["chunks"] if c["topic"] == topic and len(c["text"]) > 300] or \
           [c for c in ix["chunks"] if c["topic"] == topic]
    if not pool: return None
    ctx = random.sample(pool, min(3, len(pool)))
    if not _have_llm():
        t = ctx[0]["text"][:90].replace("\n", " ")
        opts = ["True", "False", "Cannot be determined", "None of the above"]
        return dict(question=f"[Demo] Based on {topic}: is this statement supported? \"{t}...\"",
                    options=opts, correct_idx=0, explanation="Demo mode explanation.", topic=topic)
    cfg = config.COURSES[course]
    prompt = (
        f"Create ONE multiple-choice question for a PharmD student on the topic '{topic}' of the course "
        f"{cfg['title']}, based ONLY on the material below. 4 options, exactly one correct, plausible distractors, "
        "student level, answerable without images. Respond as JSON: "
        '{"question": str, "options": [4 strings], "correct_idx": 0-3, "explanation": "2-3 sentence explanation"}.\n'
        + ("Do not repeat these questions:\n- " + "\n- ".join(avoid) + "\n" if avoid else "")
        + "\nMATERIAL:\n" + "\n\n".join(c["text"] for c in ctx))
    for _ in range(2):
        try:
            d = json.loads(_chat([{"role": "user", "content": prompt}], json_mode=True, temperature=0.8, max_tokens=500))
            if len(d["options"]) == 4 and 0 <= int(d["correct_idx"]) <= 3:
                d["correct_idx"] = int(d["correct_idx"]); d["topic"] = topic
                return d
        except Exception:
            continue
    return None

def recommendations(course, stats):
    """stats is anonymised aggregate data (no student names)."""
    if not _have_llm():
        weak = ", ".join(t["topic"] for t in stats["weak_topics"][:3]) or "none yet"
        return ("[Demo mode, no OpenAI key]\n1. Review the weakest topics: " + weak +
                ".\n2. Add worked examples and short practice sets for them.\n3. Add a recap quiz at the start of each lecture.")
    cfg = config.COURSES[course]
    prompt = (
        f"You advise the instructor of {cfg['title']} ({cfg['code']}). Below is anonymised analytics from the AI tutor "
        "used by students. Write practical recommendations to improve the course: (1) what to re-teach or clarify, "
        "(2) content/material improvements, (3) activities or assessments to add, (4) students/groups needing follow-up "
        "(describe by pattern, not by name). Be specific, cite the topics and numbers, max 350 words, plain text with "
        "numbered points, no markdown symbols.\n\nDATA:\n" + json.dumps(stats, ensure_ascii=False)[:9000])
    return _chat([{"role": "user", "content": prompt}], temperature=0.4, max_tokens=700)


# ---------------------------------------------------------------- theory questions
def _sample_ctx(course, topic, n=3):
    ix = load(course)
    pool = [c for c in ix["chunks"] if c["topic"] == topic and len(c["text"]) > 300] or \
           [c for c in ix["chunks"] if c["topic"] == topic]
    return random.sample(pool, min(n, len(pool))) if pool else []

def _json(prompt, max_tokens=900, temperature=0.7):
    for _ in range(2):
        try:
            return json.loads(_chat([{"role": "user", "content": prompt}], json_mode=True,
                                    temperature=temperature, max_tokens=max_tokens))
        except Exception:
            continue
    return None

def make_theory(course, topic, avoid=()):
    ctx = _sample_ctx(course, topic)
    if not ctx: return None
    if not _have_llm():
        t = re.sub(r"\s+", " ", ctx[0]["text"])[:300]
        return dict(question=f"[Demo] Explain the main idea of this part of {topic} in your own words.",
                    model_answer=t, topic=topic)
    cfg = config.COURSES[course]
    prompt = (f"Create ONE short-answer (theory) question for a PharmD student on the topic '{topic}' of the course "
              f"{cfg['title']}, based ONLY on the material below. It must be answerable in 3-6 sentences without images. "
              'Respond as JSON: {"question": str, "model_answer": "a complete model answer of 3-6 sentences"}.\n'
              + ("Do not repeat these questions:\n- " + "\n- ".join(avoid) + "\n" if avoid else "")
              + "\nMATERIAL:\n" + "\n\n".join(c["text"] for c in ctx))
    d = _json(prompt, 600, 0.8)
    if d and d.get("question") and d.get("model_answer"):
        return dict(question=str(d["question"]), model_answer=str(d["model_answer"]), topic=topic)
    return None

def grade_theory(question, model_answer, student_answer):
    """Returns dict(score 0-100, feedback)."""
    if _have_llm():
        prompt = ("You are a fair, encouraging examiner. Grade the STUDENT ANSWER against the MODEL ANSWER for the question. "
                  "Give credit for correct ideas even if worded differently; do not require exact wording. "
                  'Respond as JSON: {"score": integer 0-100, "feedback": "2-3 sentences: what was right, what is missing"}. '
                  "Reply in the language of the student's answer.\n"
                  f"QUESTION: {question}\nMODEL ANSWER: {model_answer}\nSTUDENT ANSWER: {student_answer}")
        d = _json(prompt, 300, 0.2)
        if d and "score" in d:
            try: return dict(score=max(0, min(100, int(float(d["score"])))), feedback=str(d.get("feedback", "")))
            except Exception: pass
    a, b = set(tokenize(model_answer)), set(tokenize(student_answer))
    a = {w for w in a if len(w) > 3}
    overlap = len(a & b) / max(1, len(a))
    score = int(min(100, overlap * 140))
    return dict(score=score, feedback="Automatic estimate based on key terms (AI grading unavailable). Compare with the model answer.")

def make_bank(course, topic_list, kind, n):
    """Teacher question sheet: kind 'mcq' or 'theory'. Returns list of dicts with answers."""
    cfg = config.COURSES[course]
    ctx = []
    for t in topic_list: ctx += _sample_ctx(course, t, 2)
    random.shuffle(ctx); ctx = ctx[:10]
    if not _have_llm():
        out = []
        for i in range(n):
            c = ctx[i % len(ctx)]
            if kind == "mcq":
                out.append(dict(topic=c["topic"], question=f"[Demo] Question {i+1} on {c['topic']}", options=["A1", "B1", "C1", "D1"],
                                correct_idx=0, explanation="Demo."))
            else:
                out.append(dict(topic=c["topic"], question=f"[Demo] Question {i+1}: explain a key idea from {c['topic']}",
                                model_answer=re.sub(r"\s+", " ", c["text"])[:250]))
        return out
    shape = ('{"questions":[{"topic": str, "question": str, "options": [4 strings], "correct_idx": 0-3, "explanation": str}]}'
             if kind == "mcq" else '{"questions":[{"topic": str, "question": str, "model_answer": "3-6 sentence model answer"}]}')
    label = "multiple-choice (4 options, one correct)" if kind == "mcq" else "short-answer theory"
    labels = "\n".join(f"[{c['topic']}]\n{c['text']}" for c in ctx)
    prompt = (f"Write exactly {n} different {label} exam questions for the course {cfg['title']} ({cfg['code']}), PharmD level, "
              f"based ONLY on the material below, spread over the topics, no images needed. Respond as JSON: {shape}. "
              "Use the topic name given in brackets for each question.\n\nMATERIAL:\n" + labels)
    d = _json(prompt, 3500, 0.7)
    qs = (d or {}).get("questions") or []
    good = []
    for z in qs:
        if not z.get("question"): continue
        if kind == "mcq":
            try:
                if len(z["options"]) == 4 and 0 <= int(z["correct_idx"]) <= 3: z["correct_idx"] = int(z["correct_idx"]); good.append(z)
            except Exception: pass
        elif z.get("model_answer"): good.append(z)
    return good[:n]

def ensure_embeddings():
    """Build missing embeddings once (safe to call from several processes)."""
    if not _have_llm(): return
    for course in config.COURSES:
        jp, np_ = paths(course)
        if os.path.exists(np_): continue
        lock = np_ + ".lock"
        try: fd = os.open(lock, os.O_CREAT | os.O_EXCL); os.close(fd)
        except FileExistsError: continue
        try: build_embeddings(course)
        except Exception as e: print("embedding build failed:", e)
        finally:
            try: os.remove(lock)
            except OSError: pass
